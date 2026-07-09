#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Append postscript to the combined crisis paper docx."""

from docx import Document
from docx.shared import Pt, Cm
from docx.oxml.ns import qn

from add_part3_abstract import set_run_font, style_para

DOCX = "/workspace/1997亚洲金融危机.docx"

POSTSCRIPT_TITLE = "后记"

POSTSCRIPT = (
    "转眼，又一个学期走到了尾声。"
    "回想大学这几年，前几学期我们几乎没有接触系统性的金融专业课，"
    "对“金融学”三个字多少还是有些距离感——知道它与我们未来密切相关，"
    "却很难真正走进去。这学期修读刘老师的金融学课程，于我而言不只是一门"
    "课的结束，更像是一扇门被推开：老师上课深入浅出，把复杂的机制讲得"
    "清楚有趣，案例与理论穿插之间干货满满，让我第一次觉得，金融学不是"
    "远在天边的术语，而是可以用来理解真实世界变化的工具。"
    "本篇论文写作的过程，比我想象中要长得多，也比我预想的更有收获。"
    "为了把1997年亚洲金融危机与欧元区主权债务危机的机制与对照写清楚，"
    "我反复阅读文献、核对数据，许多理解是在课堂之外一点一点沉下来的——"
    "比如繁荣期脆弱性如何被低估，局部冲击又如何沿银行间和跨境融资网络"
    "扩散；又比如不同制度约束下，最后贷款人与问题机构处置为何呈现出"
    "截然不同的政策空间。有时为了一个出处的核对看到深夜，有时读完一篇"
    "经典文献，忽然回过味来老师课上某句话的分量，那种“原来如此”的"
    "豁然开朗，大概就是这个学期留给我的最踏实的一份收获。"
    "课程即将结束，心里多少有些不舍。非常荣幸能在本学期跟随刘老师学习，"
    "感谢老师一学期以来的悉心讲授。也衷心祝愿老师身体健康，工作顺利！"
)


def main():
    doc = Document(DOCX)

    # Remove existing postscript if re-running
    body = doc.element.body
    start = None
    for i, para in enumerate(doc.paragraphs):
        if para.text.strip() == POSTSCRIPT_TITLE:
            start = i
            break
    if start is not None:
        for para in list(doc.paragraphs)[start:]:
            body.remove(para._element)

    title = doc.add_paragraph(POSTSCRIPT_TITLE)
    style_para(title, heading=True)
    set_run_font(title.runs[0], size=12, bold=True)

    body_para = doc.add_paragraph(POSTSCRIPT)
    style_para(body_para)
    body_para.paragraph_format.first_line_indent = Cm(0.74)

    doc.save(DOCX)
    print("Appended postscript to", DOCX)


if __name__ == "__main__":
    main()
