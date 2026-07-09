#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Add abstract and Part III comparative section to combined crisis paper."""

from copy import deepcopy
from docx import Document
from docx.shared import Pt, Cm
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = "/home/ubuntu/.cursor/projects/workspace/uploads/1997_______3cde.docx"
OUT_DESKTOP = "/home/ubuntu/Desktop/1997亚洲金融危机.docx"
OUT_WORKSPACE = "/workspace/1997亚洲金融危机.docx"

NEW_TITLE = "1997年亚洲金融危机与欧元区主权债务危机：机制比较与政策启示"

ABSTRACT = (
    "摘要：本文在刘泽豪老师第十二讲“金融脆弱性与金融危机”的分析框架下，"
    "分别考察1997年亚洲金融危机与2010—2018年欧元区主权债务危机的脆弱性积累、"
    "爆发传导、政策应对与成因机制，并与课上重点讨论的美国1907年恐慌、"
    "大萧条与2007—2009年全球金融危机进行对照。两场危机均发生在金融自由化"
    "与跨境资本流动加速背景下：亚洲危机表现为固定汇率、短期外债与银行体系"
    "期限/货币错配相互强化；欧元区危机则体现为“统一货币—分散财政”的不完备"
    "货币联盟结构，以及主权—银行反馈环下的恐慌自我实现。政策层面，二者均凸显"
    "最后贷款人、存款担保/银行重组与宏观政策两难，但亚洲更多依赖IMF牵头救助"
    "与各国差异化路径，欧元区则通过EFSM/ESM、OMT与QE逐步补建区域金融稳定架构。"
    "跨案例比较表明，金融危机虽形态各异，却共享“脆弱性积累—触发事件—挤兑/重定价—"
    "实体衰退”链条；理解危机的关键在于识别制度约束如何放大金融自由化的激励扭曲，"
    "而非将冲击简单归因于投机攻击或单一政策失误。"
)

KEYWORDS = "关键词：金融脆弱性；金融自由化；货币危机；银行危机；最后贷款人；传染"

PART3 = [
    ("第三部分  跨案例比较与总结", True),
    ("三、跨案例比较：两场危机与课上经典危机的异同", True),
    ("3.1 分析框架：从金融脆弱性到系统性危机", False),
    (
        "第十二讲将金融脆弱性界定为融资体系中风险积聚的状态，并强调金融自由化"
        "在提升效率的同时可能激化固有脆弱性（Demirgüç-Kunt和Detragiache，1999）。"
        "利率自由化带来的逆向选择、存贷利差收窄与资产价格波动，业务扩张使银行与"
        "房地产/证券周期绑定，资本流动放松则放大固定汇率或不完备货币联盟下的外部"
        "失衡。本文所讨论的两场危机，以及课上的1907年恐慌、大萧条与2007—2009年"
        "全球金融危机（GFC），都可以在这一框架下理解为：繁荣期风险被低估、约束被"
        "弱化，冲击到来后通过挤兑、重定价或去杠杆迅速释放。",
        False,
    ),
    ("3.2 触发机制与危机类型：挤兑、货币与影子银行", False),
    (
        "1907年恐慌的触发点是United Copper囤购失败及Knickerbocker Trust倒闭，"
        "典型表现为信任崩溃与银行挤兑（Diamond和Dybvig，1983）；J.P.Morgan的私人"
        "协调与财政部应对不足，直接推动了美联储的建立。大萧条兼具股市崩盘、"
        "多轮银行恐慌与金本位约束下的通缩传导（Eichengreen，2002；Bernanke，1983），"
        "货币收缩与信贷中介成本上升共同放大了衰退。2007—2009年GFC则主要发生在"
        "影子银行体系：次贷基本面恶化引发ABX价格下跌，回购与ABCP市场“挤兑”，"
        "LIBOR-OIS利差飙升显示银行间信用风险重定价（Gorton和Metrick，2012）。"
        "1997年亚洲危机以泰国房地产/金融公司问题触发本币失守，并快速演变为"
        "货币—银行—实体三重危机；2010年欧元区危机由希腊财政数据修正触发，"
        "但深层是利率收敛下的信贷扩张、跨境银行敞口与竞争力分化，危机类型呈现"
        "“主权—银行”耦合与债务可持续性恐慌。",
        False,
    ),
    ("3.3 传染与反馈环：区域同步重定价的共性", False),
    (
        "Kindleberger（1978）强调恐慌在开放条件下的多米诺式传染。亚洲危机中，"
        "国际投资者按“新兴亚洲”整体削减敞口，泰铢贬值通过贸易与预期渠道波及其他"
        "东盟经济体；韩国则因短期外债滚续中断而遭遇批发融资“挤兑”。欧元区危机的"
        "传染更体现为评级下调的顺周期性、回购抵押品折扣上升与跨境银行共同持有"
        "PIIGS主权债的“互相踩踏”，使希腊问题迅速外溢至爱尔兰、葡萄牙、西班牙乃至"
        "意大利。与1907年、1930年代美国银行恐慌相似，信息不完全条件下市场往往"
        "“按类别”而非“按基本面”抛售；与GFC相似，银行间市场冻结是危机深化的关键"
        "加速器。差异在于：亚洲与欧元区的冲击主要沿跨境资本与主权/银行资产负债表"
        "传导，而GFC还叠加了复杂证券化链条上的抵押品质量重估。",
        False,
    ),
    ("3.4 政策应对：最后贷款人、重组与道德风险", False),
    (
        "1907年缺乏制度化最后贷款人，危机后以联邦储备体系作为制度回应；大萧条"
        "期间美联储未能有效履行LOLR职能，1933年存款担保与银行假日才逐步稳定体系。"
        "GFC中，美联储与财政部通过非常规流动性工具、问题机构处置与存款保险稳定"
        "预期，但“大而不倒”与道德风险争议持续存在。亚洲危机的政策组合是放弃不可"
        "持续钉住、IMF牵头官方融资、关闭问题机构并扩展存款担保；马来西亚资本管制"
        "提供了与IMF模板并行的实验。欧元区则经历“先紧缩、后承诺”的演进："
        "EFSF/ESM提供区域救助，2012年德拉吉“Whatever it takes”与OMT以承诺型"
        "最后贷款人稳定预期，2015年后QE进一步压低主权利差，银行业联盟试图切断"
        "主权—银行螺旋。共性是：仅货币宽松不足以修复受损的银行中介（Bernanke，1983）；"
        "但若缺乏透明重组与硬约束，公共救助又会积累新的道德风险。",
        False,
    ),
    ("3.5 两场危机的同异对照", False),
    (
        "相同点：（1）危机前均存在金融自由化背景下的信贷扩张与资产价格上升；"
        "（2）都出现货币/主权与银行部门的双向反馈；（3）都伴随国际投资者同步"
        "重定价与区域传染；（4）政策均在“稳定金融”与“避免道德风险”之间艰难"
        "权衡。不同点：（1）制度环境：亚洲为中等收入开放经济体+固定汇率，欧元区"
        "为高收入货币联盟内部的主权债务问题；（2）核心错配：亚洲突出短期外债与"
        "货币错配，欧元区突出财政分散与缺乏共同最后贷款人；（3）冲击来源：亚洲"
        "先有实体/金融部门问题再冲击汇率，欧元区先有主权信用重估再冲击银行；"
        "（4）政策工具：亚洲以IMF双边/多边救助为主，欧元区逐步内生构建ESM/OMT/QE"
        "等区域工具。",
        False,
    ),
    ("3.6 总结与启示", False),
    (
        "综合五类危机案例，可以得到三点启示。第一，金融稳定不是增长的自动副产品，"
        "金融自由化必须与监管能力、信息披露和危机处置框架同步推进。第二，固定汇率、"
        "不完备货币联盟或影子银行体系中的隐性担保，都会在高杠杆环境下放大顺周期"
        "行为；危机治理的关键是打破“贬值/违约—银行受损—资本外流”的反馈环。"
        "第三，最后贷款人与存款/融资稳定机制是遏制挤兑的必要条件，但其设计必须"
        "与事后重组、损失分担和财政约束相配套，否则只能以时间换空间，难以根除"
        "脆弱性。对新兴经济体而言，亚洲危机提示应谨慎管理短期外债与汇率制度；"
        "对货币联盟建设而言，欧元区危机提示仅有货币一体化而缺少财政与金融联盟，"
        "可能在冲击下付出更长的调整成本。",
        False,
    ),
]


def set_run_font(run, name_cn="宋体", name_en="Times New Roman", size=10.5, bold=False):
    run.font.name = name_en
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name_cn)
    run.font.size = Pt(size)
    run.bold = bold


def style_para(para, *, heading=False, abstract=False, keywords=False):
    pf = para.paragraph_format
    pf.line_spacing = 1.5
    pf.space_after = Pt(0)
    if heading:
        pf.space_before = Pt(12)
        pf.space_after = Pt(6)
        pf.first_line_indent = Cm(0)
        if para.runs:
            set_run_font(para.runs[0], size=12, bold=True)
    elif keywords:
        pf.first_line_indent = Cm(0)
        if para.runs:
            set_run_font(para.runs[0], size=10.5)
    elif abstract:
        pf.first_line_indent = Cm(0.74)
        if para.runs:
            set_run_font(para.runs[0], size=10.5)
    else:
        pf.first_line_indent = Cm(0.74)
        if para.runs:
            set_run_font(para.runs[0], size=10.5)


def insert_paragraph_before(paragraph, text="", style=None):
    new_p = OxmlElement("w:p")
    paragraph._p.addprevious(new_p)
    from docx.text.paragraph import Paragraph

    new_para = Paragraph(new_p, paragraph._parent)
    if style:
        new_para.style = style
    if text:
        run = new_para.add_run(text)
        set_run_font(run, size=12, bold=True)
    return new_para


def insert_paragraph_after(paragraph, text="", style=None):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    from docx.text.paragraph import Paragraph

    new_para = Paragraph(new_p, paragraph._parent)
    if style:
        new_para.style = style
    if text:
        run = new_para.add_run(text)
        set_run_font(run)
    return new_para


def find_para_index(doc, predicate):
    for i, p in enumerate(doc.paragraphs):
        if predicate(p.text.strip()):
            return i
    return -1


def main():
    doc = Document(SRC)

    # Update title
    if doc.paragraphs:
        p0 = doc.paragraphs[0]
        if p0.runs:
            p0.runs[0].text = NEW_TITLE
            for r in p0.runs[1:]:
                r.text = ""
        else:
            p0.text = NEW_TITLE
        p0.alignment = 1
        set_run_font(p0.runs[0], size=16, bold=True)

    # Insert abstract + keywords after title (order: abstract then keywords)
    title_p = doc.paragraphs[0]
    abs_p = insert_paragraph_after(title_p, ABSTRACT)
    style_para(abs_p, abstract=True)
    kw_p = insert_paragraph_after(abs_p, KEYWORDS)
    style_para(kw_p, keywords=True)

    # Part labels before each major block
    for p in list(doc.paragraphs):
        t = p.text.strip()
        if t == "一、危机前的东亚：增长、自由化与不可持续的金融结构":
            insert_paragraph_before(p, "第一部分  1997年亚洲金融危机")
        elif t == "一、欧元区制度安排的内在张力与危机前的失衡积累":
            insert_paragraph_before(p, "第二部分  欧元区主权债务危机")

    # Insert Part III before references block (after euro summary)
    indices = [i for i, p in enumerate(doc.paragraphs) if p.text.strip() == "参考文献"]
    insert_before = indices[0] if indices else len(doc.paragraphs) - 1
    anchor = doc.paragraphs[insert_before - 1]

    chain = anchor
    for text, is_heading in PART3:
        chain = insert_paragraph_after(chain, text)
        style_para(chain, heading=is_heading)

    # Merge duplicate reference sections
    indices = [i for i, p in enumerate(doc.paragraphs) if p.text.strip() == "参考文献"]
    if len(indices) >= 2:
        second = indices[1]
        asian_refs = [p.text.strip() for p in doc.paragraphs[second + 1 :] if p.text.strip()]
        body = doc.element.body
        for p in list(doc.paragraphs)[second:]:
            body.remove(p._element)
        # collect existing euro refs
        existing = set()
        tail = None
        for p in doc.paragraphs:
            if p.text.strip() == "参考文献":
                tail = p
                continue
            if tail is not None and p.text.strip():
                existing.add(p.text.strip())
                tail = p
        if tail is None:
            tail = doc.paragraphs[-1]
        for ref in asian_refs:
            if ref in existing:
                continue
            tail = insert_paragraph_after(tail, ref)
            tail.paragraph_format.left_indent = Cm(0.74)
            tail.paragraph_format.first_line_indent = Cm(-0.74)
            tail.paragraph_format.line_spacing = 1.5
            if tail.runs:
                set_run_font(tail.runs[0], size=10.5)
            existing.add(ref)

    doc.save(OUT_DESKTOP)
    doc.save(OUT_WORKSPACE)
    print("Saved:", OUT_DESKTOP)


def replace_text(para, new_text):
    if para.runs:
        para.runs[0].text = new_text
        for r in para.runs[1:]:
            r.text = ""
    else:
        para.text = new_text


if __name__ == "__main__":
    main()
